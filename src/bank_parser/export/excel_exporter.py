"""Formatted .xlsx writer: a Transactions sheet (every row, every statement)
and a Summary sheet (one row per statement, color-coded reconciliation
status) -- doc section 27's main screen, on paper."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import IO, TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from bank_parser.models.statement import BankStatement

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_OK_FILL = PatternFill("solid", fgColor="C6EFCE")
_REVIEW_FILL = PatternFill("solid", fgColor="FFC7CE")

_TXN_HEADERS = [
    "source_file",
    "transaction_date",
    "posting_date",
    "description",
    "debit",
    "credit",
    "amount",
    "balance",
    "confidence",
]

_SUMMARY_HEADERS = [
    "source_file",
    "institution",
    "period_start",
    "period_end",
    "transactions",
    "debit_extracted",
    "debit_printed",
    "credit_extracted",
    "credit_printed",
    "balance_checks_passed",
    "score",
    "reconciled",
    "warnings",
]


def _style_header_row(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autosize(ws: Worksheet, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 2, max_width)


def _write_transactions(ws: Worksheet, statements: list["BankStatement"]) -> None:
    ws.append(_TXN_HEADERS)
    _style_header_row(ws)

    for statement in statements:
        rows = sorted(statement.transactions, key=lambda t: t.transaction_date or date.min)
        for txn in rows:
            ws.append(
                [
                    statement.source_filename,
                    txn.transaction_date,
                    txn.posting_date,
                    txn.description,
                    float(txn.debit) if txn.debit is not None else None,
                    float(txn.credit) if txn.credit is not None else None,
                    float(txn.amount) if txn.amount is not None else None,
                    float(txn.balance) if txn.balance is not None else None,
                    round(txn.confidence, 2),
                ]
            )

    date_columns = ("B", "C")
    money_columns = ("E", "F", "G", "H")
    for row in ws.iter_rows(min_row=2):
        for col in date_columns:
            row[ord(col) - ord("A")].number_format = "yyyy-mm-dd"
        for col in money_columns:
            row[ord(col) - ord("A")].number_format = "#,##0.00"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def _write_summary(ws: Worksheet, statements: list["BankStatement"]) -> None:
    ws.append(_SUMMARY_HEADERS)
    _style_header_row(ws)

    for statement in statements:
        v = statement.validation
        d = statement.diagnostics
        ws.append(
            [
                statement.source_filename,
                statement.institution,
                statement.period_start,
                statement.period_end,
                len(statement.transactions),
                float(d.get("extracted_debit_sum") or 0),
                d.get("printed_totals", {}).get("debit_total"),
                float(d.get("extracted_credit_sum") or 0),
                d.get("printed_totals", {}).get("credit_total"),
                f"{v.running_balance_passes}/{v.running_balance_checks}",
                round(v.score, 3),
                "OK" if v.is_reliable else "REVIEW",
                "; ".join(v.warnings) if v.warnings else "",
            ]
        )

    status_col = _SUMMARY_HEADERS.index("reconciled") + 1
    for row in ws.iter_rows(min_row=2):
        cell = row[status_col - 1]
        cell.fill = _OK_FILL if cell.value == "OK" else _REVIEW_FILL

    for row in ws.iter_rows(min_row=2):
        row[_SUMMARY_HEADERS.index("period_start")].number_format = "yyyy-mm-dd"
        row[_SUMMARY_HEADERS.index("period_end")].number_format = "yyyy-mm-dd"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def write_excel(statements: list["BankStatement"], target: str | Path | IO[bytes]) -> None:
    """`target` is a filesystem path, or a writable file-like object (e.g. an
    in-memory `io.BytesIO` for a UI download button, never touching disk)."""
    if isinstance(target, (str, Path)):
        target = Path(target)
        target.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    txn_sheet = wb.active
    txn_sheet.title = "Transactions"
    _write_transactions(txn_sheet, statements)

    summary_sheet = wb.create_sheet("Summary")
    _write_summary(summary_sheet, statements)

    wb.save(target)
