"""Write extracted statements to a formatted Excel workbook."""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import models
from .models import StatementResult, Transaction

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MISMATCH_FILL = PatternFill("solid", fgColor="F8CBAD")
_OK_FILL = PatternFill("solid", fgColor="C6E0B4")
_MONEY_FMT = "#,##0.00"
_DATE_FMT = "yyyy-mm-dd"

_HEADERS = [
    "source_file", "statement_type", "fecha_operacion", "fecha_cargo",
    "description", "details", "money_out", "money_in", "balance", "category",
]


def write_workbook(results: List[StatementResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_transactions(wb.active, results)
    _write_summary(wb.create_sheet("Summary"), results)
    wb.save(output_path)


def _all_rows(results: List[StatementResult]) -> List[Transaction]:
    rows: List[Transaction] = []
    for r in results:
        rows.extend(r.transactions)
        rows.extend(r.installments)
    rows.sort(key=lambda t: (t.source_file, t.fecha_operacion or date.min))
    return rows


def _write_transactions(ws, results: List[StatementResult]) -> None:
    ws.title = "Transactions"
    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for tx in _all_rows(results):
        row = tx.to_row()
        ws.append([row[h] for h in _HEADERS])

    _format_columns(ws, n_header=1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{ws.max_row}"


def _format_columns(ws, n_header: int) -> None:
    col_idx = {h: i + 1 for i, h in enumerate(_HEADERS)}
    for r in range(n_header + 1, ws.max_row + 1):
        for name in ("fecha_operacion", "fecha_cargo"):
            ws.cell(r, col_idx[name]).number_format = _DATE_FMT
        for name in ("money_out", "money_in", "balance"):
            ws.cell(r, col_idx[name]).number_format = _MONEY_FMT
    widths = {
        "source_file": 30, "statement_type": 18, "fecha_operacion": 14,
        "fecha_cargo": 14, "description": 40, "details": 55,
        "money_out": 13, "money_in": 13, "balance": 14, "category": 15,
    }
    for name, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx[name])].width = w


def _write_summary(ws, results: List[StatementResult]) -> None:
    # "spending_*" = clean output (internal transfers removed). "bbva_total_*" = the
    # figures BBVA prints, which include the net-zero PROMOCION/TRASPASO pairs.
    # spending_out + internal_excluded == bbva_total_cargos (the reconciliation check).
    headers = [
        "source_file", "type", "period_start", "period_end", "txns",
        "spending_out", "spending_in", "internal_excluded",
        "installments", "installment_total",
        "bbva_total_cargos", "bbva_total_abonos", "reconciled",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for r in results:
        out = sum(t.money_out for t in r.transactions if t.money_out) or 0.0
        inc = sum(t.money_in for t in r.transactions if t.money_in) or 0.0
        inst_total = sum(t.money_out for t in r.installments if t.money_out) or 0.0
        # Amount removed as internal transfers (equal on the cargo & abono sides).
        internal = (r.extracted_cargos or 0.0) - out
        start = r.period[0] if r.period else None
        end = r.period[1] if r.period else None
        ws.append([
            r.source_file, r.statement_type, start, end, len(r.transactions),
            round(out, 2), round(inc, 2), round(internal, 2),
            len(r.installments), round(inst_total, 2),
            r.printed_cargos, r.printed_abonos, "OK" if r.reconciled else "MISMATCH",
        ])
        cell = ws.cell(ws.max_row, len(headers))
        cell.fill = _OK_FILL if r.reconciled else _MISMATCH_FILL

    # Number/date formatting for the statement block.
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 3).number_format = _DATE_FMT
        ws.cell(row, 4).number_format = _DATE_FMT
        for col in (6, 7, 8, 10, 11, 12):
            ws.cell(row, col).number_format = _MONEY_FMT
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["A"].width = 30

    _write_category_block(ws, results, start_row=ws.max_row + 3)


def _write_category_block(ws, results: List[StatementResult], start_row: int) -> None:
    totals = defaultdict(lambda: [0.0, 0.0])  # category -> [out, in]
    for r in results:
        for t in list(r.transactions) + list(r.installments):
            cat = t.category or "unknown"
            totals[cat][0] += t.money_out or 0.0
            totals[cat][1] += t.money_in or 0.0

    title = ws.cell(start_row, 1, "Category totals")
    title.font = Font(bold=True)
    hdr = ["category", "money_out", "money_in"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(start_row + 1, j, h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
    for k, (cat, (out, inc)) in enumerate(sorted(totals.items()), start=start_row + 2):
        ws.cell(k, 1, cat)
        ws.cell(k, 2, round(out, 2)).number_format = _MONEY_FMT
        ws.cell(k, 3, round(inc, 2)).number_format = _MONEY_FMT
