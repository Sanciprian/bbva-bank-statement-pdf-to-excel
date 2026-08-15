from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from bank_parser.models.statement import BankStatement
from bank_parser.models.transaction import Transaction
from bank_parser.models.validation import StatementValidation


def _statement():
    txn = Transaction(
        transaction_date=date(2025, 10, 10),
        posting_date=date(2025, 10, 10),
        description="SPEI RECIBIDO",
        debit=None,
        credit=Decimal("1541.00"),
        amount=Decimal("1541.00"),
        balance=Decimal("10505.13"),
        reference=None,
        source_page=1,
        source_bbox=None,
        confidence=0.9,
    )
    return BankStatement(
        source_filename="synthetic.pdf",
        institution="BBVA",
        account_number_masked=None,
        currency=None,
        period_start=date(2025, 10, 9),
        period_end=date(2025, 11, 8),
        opening_balance=None,
        closing_balance=Decimal("10505.13"),
        transactions=[txn],
        confidence=0.9,
        validation=StatementValidation(score=0.9),
    )


def test_to_dataframe_round_trip():
    df = _statement().to_dataframe()
    assert len(df) == 1
    row = df.iloc[0]
    assert row["credit"] == 1541.00
    assert row["balance"] == 10505.13
    assert row["source_file"] == "synthetic.pdf"


def test_to_csv_writes_numeric_values(tmp_path):
    path = tmp_path / "out.csv"
    _statement().to_csv(path)
    content = path.read_text()
    assert "1541.0" in content
    assert "10505.13" in content


def test_to_excel_writes_numeric_not_string_cells(tmp_path):
    path = tmp_path / "out.xlsx"
    _statement().to_excel(path)

    wb = load_workbook(path)
    ws = wb["Transactions"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    values = dict(zip(header, row))

    assert isinstance(values["credit"], (int, float))
    assert values["credit"] == 1541.00
    assert isinstance(values["balance"], (int, float))
    assert values["balance"] == 10505.13
